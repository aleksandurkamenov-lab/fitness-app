#!/usr/bin/env python3
"""Превръща треньорския Excel (Upper/Lower блокове един до друг) в JSON за приложението.

Структура на листа:
  - на ред 6 са имената на тренировките, всяка в свой блок от колони;
  - всяко упражнение е блок от 7 реда: заглавие, КГ/Повт, най-добри резултати, седмици 1-4;
  - серия i има три колони: Target, КГ, Повт.

Употреба: python3 tools/xlsx_to_json.py файл.xlsx FitnessApp/Resources/Meso22.json
"""
import json
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string as col_idx

COMMENT_LABEL = "Коментар по упражнението"
SKIP_COMMENTS = {"Link", "Най - добри резултати до сега -->"}

MUSCLE_KEYWORDS = [
    ("calves", ["calves", "calf"]),
    ("glutes", ["hip thrust", "glute"]),
    ("hamstrings", ["rdl", "leg curl", "hamstring", "deadlift"]),
    ("quads", ["leg press", "squat", "leg extension", "lunge", "adductor"]),
    ("abs", ["crunch", "abs", "leg raise"]),
    ("biceps", ["curl", "biceps"]),
    ("triceps", ["tricep", "pushdown", "extension", "skull"]),
    ("shoulders", ["lateral", "overhead", "shoulder", "reverce", "reverse pec", "face pull", "rear delt"]),
    ("back", ["pulldown", "row", "pull-up", "pullup", "shrug", "lat "]),
    ("chest", ["bench", "chest", "pec", "fly", "push"]),
]


def guess_muscle(name: str) -> str:
    lowered = name.lower()
    for muscle, words in MUSCLE_KEYWORDS:
        if any(word in lowered for word in words):
            return muscle
    return "chest"


def parse_range(value):
    if not isinstance(value, str):
        return None
    match = re.match(r"^\s*(\d+)\s*-\s*(\d+)\s*$", value)
    return (int(match.group(1)), int(match.group(2))) if match else None


def number(value):
    return float(value) if isinstance(value, (int, float)) else None


def convert(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    max_col = ws.max_column

    starts = [c for c in range(1, max_col + 1) if isinstance(ws.cell(6, c).value, str) and ws.cell(6, c).value.strip()]
    workouts = []
    for start in starts:
        exercises = []
        for row in range(7, ws.max_row + 1):
            name = ws.cell(row, start).value
            if not (isinstance(name, str) and name.strip()):
                continue
            if ws.cell(row, start + 3).value != COMMENT_LABEL:
                continue

            sets = []
            offset = 5
            while offset + 2 < max_col - start + 1:
                target = parse_range(ws.cell(row + 3, start + offset).value)
                if target is None:
                    break
                sets.append({
                    "repMin": target[0],
                    "repMax": target[1],
                    "bestWeight": number(ws.cell(row + 2, start + offset + 1).value),
                    "bestReps": number(ws.cell(row + 2, start + offset + 2).value),
                })
                offset += 3

            notes = []
            for r in range(row + 1, row + 7):
                text = ws.cell(r, start + 3).value
                if isinstance(text, str) and text.strip() and text.strip() not in SKIP_COMMENTS:
                    notes.append(text.strip())

            exercises.append({
                "name": name.strip(),
                "muscle": guess_muscle(name),
                "note": " · ".join(notes),
                "sets": sets,
            })
        workouts.append({"name": ws.cell(6, start).value.strip(), "exercises": exercises})
    return {"name": "Примерен мезоцикъл", "weeks": 4, "workouts": workouts}


if __name__ == "__main__":
    data = convert(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as handle:
        json.dump(data, handle, ensure_ascii=False, indent=2)
    total = sum(len(w["exercises"]) for w in data["workouts"])
    print(f"{len(data['workouts'])} тренировки, {total} упражнения")
